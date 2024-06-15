import pyttsx3
import os
import PyPDF2
from docx import Document
import openpyxl

def speak_text(text):
    # Inicializa o objeto da engine de síntese de voz
    engine = pyttsx3.init()
    
    # Configura a voz para português brasileiro, se disponível
    brazilian_voice = None
    for voice in engine.getProperty('voices'):
        languages = getattr(voice, 'languages', [])
        if languages:
            for language in languages:
                if "brazil" in language.lower():
                    brazilian_voice = voice.id
                    break
        if brazilian_voice:
            break
    if brazilian_voice:
        engine.setProperty('voice', brazilian_voice)
    else:
        print("Voz em português brasileiro não encontrada. Usando a voz padrão.")
    
    # Configura a velocidade de fala para um valor que soe natural
    engine.setProperty('rate', 140)  # Ajuste conforme preferência
    
    # Sintetiza e reproduz o texto
    engine.say(text)
    engine.runAndWait()

def read_text_from_txt(file_path):
    try:
        # Lê o conteúdo do arquivo .txt
        with open(file_path, 'r') as file:
            text = file.read()
        return text
    except FileNotFoundError:
        print("Arquivo não encontrado.")
        return None

def read_text_from_pdf(file_path):
    try:
        # Lê o conteúdo do arquivo PDF
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfFileReader(file)
            text = ""
            for page_num in range(reader.numPages):
                page = reader.getPage(page_num)
                text += page.extractText()
        return text
    except FileNotFoundError:
        print("Arquivo não encontrado.")
        return None

def read_text_from_docx(file_path):
    try:
        # Lê o conteúdo do arquivo .docx
        doc = Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text
        return text
    except FileNotFoundError:
        print("Arquivo não encontrado.")
        return None

def read_text_from_excel(file_path):
    try:
        # Lê o conteúdo do arquivo Excel
        wb = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        text += str(cell.value) + " "
        return text
    except FileNotFoundError:
        print("Arquivo não encontrado.")
        return None

if __name__ == "__main__":
    # Solicita ao usuário o nome do arquivo e sua extensão
    file_name = input("Digite o nome do arquivo (com extensão): ")
    
    # Verifica a extensão do arquivo e lê o conteúdo
    file_extension = os.path.splitext(file_name)[1].lower()
    if file_extension == '.txt':
        text = read_text_from_txt(file_name)
    elif file_extension == '.pdf':
        text = read_text_from_pdf(file_name)
    elif file_extension == '.docx':
        text = read_text_from_docx(file_name)
    elif file_extension == '.xlsx':
        text = read_text_from_excel(file_name)
    else:
        print("Formato de arquivo não suportado.")
        text = None
    
    # Se o texto foi lido com sucesso, converte em voz
    if text:
        speak_text(text)